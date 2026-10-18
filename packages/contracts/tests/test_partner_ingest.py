"""Partner ingestion against the injected-fault manifest of the SYNTHETIC messy dataset.

The manifest is written by scripts/make_messy_partner_dataset.py, which does not import the code under test.
"""

from __future__ import annotations

import csv
import json
import shutil
from datetime import date
from pathlib import Path

import pytest
from contracts.partner import PartnerIngestAborted, ingest_path, render_markdown
from contracts.partner.hl7_onsite import convert_messages, hmac_pseudonym
from contracts.partner.ingest import FATAL_CODES, ROW_CODES
from contracts.partner.tables import TABLES

ROOT = Path(__file__).resolve().parents[3]
FIX = ROOT / "data" / "fixtures" / "partner_messy"
MANIFEST = json.loads((FIX / "INJECTED_FAULTS.json").read_text(encoding="utf-8"))


def _expected():
    return sorted(
        (f["table"], f["row"], f["code"], f["field"]) for f in MANIFEST["expected_rejections"]
    )


def _actual(result):
    return sorted(
        (r["table"], r["row"], r["code"], r["field"]) for r in result.report["rejections"]
    )


@pytest.fixture(scope="module")
def csv_result():
    return ingest_path(FIX)


def test_every_injected_fault_is_rejected_with_its_code_and_nothing_else_is(csv_result):
    assert _actual(csv_result) == _expected()


def test_excel_and_csv_give_identical_results(csv_result):
    xl = ingest_path(FIX / "partner_messy.xlsx")
    assert _actual(xl) == _actual(csv_result)
    assert xl.report["per_table"] == csv_result.report["per_table"]
    assert xl.tables["lab_observations"] == csv_result.tables["lab_observations"]


def test_implausible_values_are_flagged_and_withheld_not_rejected(csv_result):
    flagged = {(f["row"], f["analyte"]) for f in csv_result.report["flagged_labs"]}
    assert flagged == {(f["row"], f["analyte"]) for f in MANIFEST["expected_flagged_labs"]}
    clean_rows = {r["source_row"] for r in csv_result.tables["lab_observations"]}
    assert not ({row for row, _ in flagged} & clean_rows)


def test_duplicates_and_normalisation_are_counted(csv_result):
    assert csv_result.report["duplicates_dropped"] == MANIFEST["expected_duplicates_dropped"]
    assert csv_result.report["normalised_cells"] >= MANIFEST["expected_normalised_cells_at_least"]


def test_units_are_converted_to_the_canonical_unit_not_just_accepted(csv_result):
    creat = [r for r in csv_result.tables["lab_observations"] if r["analyte"] == "creatinine"]
    assert creat and all(
        r["unit_canonical"] == "mg/dL" and 0.5 <= r["value_canonical"] <= 1.3 for r in creat
    )
    assert any(r["unit_ucum"] == "umol/L" for r in creat)  # the conversion path was exercised


def test_row_counts_reconcile_for_every_table(csv_result):
    """rows in = clean + rejected + flagged + exact duplicates dropped, and the totals match the manifest."""
    for name, v in csv_result.report["per_table"].items():
        assert (
            v["rows_clean"] + v["rows_rejected"] + v["rows_flagged"] + v["duplicates_dropped"]
            == v["rows_in"]
        ), name
        assert v["rows_in"] == MANIFEST["rows"][name]


def test_stop_rate_threshold_is_reported(csv_result):
    tight = ingest_path(FIX, stop_rate=0.01).report
    loose = ingest_path(FIX, stop_rate=0.5).report
    assert tight["exceeds_stop_rate"] is True and loose["exceeds_stop_rate"] is False
    assert "EXCEEDS" in render_markdown(tight)


def test_the_report_never_contains_a_rejected_free_text_value(csv_result):
    blob = json.dumps(csv_result.report) + render_markdown(csv_result.report)
    assert "Okafor" not in blob and "Mrs" not in blob
    assert "furlongs" not in blob and '"ten"' not in blob


def test_ingestion_is_deterministic(csv_result):
    assert ingest_path(FIX).report == csv_result.report


def test_complete_case_cohort_excludes_rejected_patients(csv_result):
    ids = set(csv_result.cohort_ids)
    rejected = {f"PTR-{i:03d}" for i in (3, 6, 9, 12, 15, 18)}
    assert not (ids & rejected) and len(ids) == csv_result.report["cohort"]["complete_case"] > 0


# --------------------------------------------------------------------------------------------- whole-file refusals
def _copy(tmp_path: Path) -> Path:
    dst = tmp_path / "x"
    shutil.copytree(FIX, dst)
    return dst


def _edit(path: Path, fn):
    with path.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
        cols = list(rows[0].keys())
    cols, rows = fn(cols, rows)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def _add_column(name, value):
    def fn(cols, rows):
        for r in rows:
            r[name] = value
        return cols + [name], rows

    return fn


def _set_cell(col, value, row=0):
    def fn(cols, rows):
        rows[row][col] = value
        return cols, rows

    return fn


@pytest.mark.parametrize(
    "table,fn,code",
    [
        ("patients", _add_column("patient_name", "Jane Doe"), "IDENTIFIER_COLUMN"),
        ("patients", _add_column("Phone", "x"), "IDENTIFIER_COLUMN"),
        ("patients", _add_column("hospital_no", "x"), "IDENTIFIER_COLUMN"),
        ("patients", _add_column("clinic_notes", "x"), "UNKNOWN_COLUMN"),
        ("patients", _set_cell("menopausal_status", "2019-03-04"), "IDENTIFIER_VALUE"),
        ("baseline_profile", _set_cell("surgery_type", "04/03/2019"), "IDENTIFIER_VALUE"),
        ("patients", _set_cell("patient_id", "08031234567"), "IDENTIFIER_VALUE"),
        ("patients", _set_cell("menopausal_status", "contact a.b@example.org"), "IDENTIFIER_VALUE"),
        ("patients", _set_cell("menopausal_status", "+234 803 123 4567"), "IDENTIFIER_VALUE"),
    ],
)
def test_a_file_with_a_direct_identifier_is_refused_as_a_whole(tmp_path, table, fn, code):
    d = _copy(tmp_path)
    _edit(d / f"{table}.csv", fn)
    with pytest.raises(PartnerIngestAborted) as exc:
        ingest_path(d)
    report = exc.value.report
    assert report["aborted"] is True and code in {f["code"] for f in report["fatal"]}
    blob = json.dumps(report) + render_markdown(report)
    for leaked in (
        "Jane Doe",
        "2019-03-04",
        "04/03/2019",
        "08031234567",
        "a.b@example.org",
        "803 123",
    ):
        assert leaked not in blob
    assert "tables" not in report  # nothing was parsed


def test_missing_table_and_missing_required_column_are_refused(tmp_path):
    d = _copy(tmp_path)
    (d / "clinical_outcomes.csv").unlink()
    with pytest.raises(PartnerIngestAborted) as exc:
        ingest_path(d)
    assert {"code": "MISSING_TABLE", "table": "clinical_outcomes"} in exc.value.report["fatal"]
    d2 = _copy(tmp_path / "again")
    _edit(
        d2 / "patients.csv",
        lambda cols, rows: (
            [c for c in cols if c != "vital_status"],
            [{k: v for k, v in r.items() if k != "vital_status"} for r in rows],
        ),
    )
    with pytest.raises(PartnerIngestAborted) as exc2:
        ingest_path(d2)
    assert {
        "code": "MISSING_COLUMN",
        "table": "patients",
        "column": "vital_status",
    } in exc2.value.report["fatal"]


def test_an_excel_date_cell_is_refused_as_an_absolute_date(tmp_path):
    from datetime import datetime

    from openpyxl import load_workbook

    p = tmp_path / "w.xlsx"
    shutil.copy(FIX / "partner_messy.xlsx", p)
    wb = load_workbook(p)
    wb["patients"]["C2"] = datetime(2019, 3, 4)
    wb.save(p)
    with pytest.raises(PartnerIngestAborted) as exc:
        ingest_path(p)
    assert any(
        f["code"] == "IDENTIFIER_VALUE" and f["kind"] == "ISO_DATE"
        for f in exc.value.report["fatal"]
    )


# --------------------------------------------------------------------------------------------- HL7 on site
def _messages():
    text = (FIX / "hl7" / "oru_messages.hl7").read_text(encoding="utf-8")
    return [m for m in text.split("\n\n") if m.strip()]


def _index():
    raw = json.loads((FIX / "hl7" / "index_dates.json").read_text(encoding="utf-8"))
    return {k: date.fromisoformat(v) for k, v in raw.items()}


def test_hl7_conversion_matches_the_manifest_and_exports_no_identifier():
    rows, rejected = convert_messages(
        _messages(), _index(), hmac_pseudonym(b"0123456789abcdef-test-only")
    )
    exp = MANIFEST["hl7"]
    assert len(rows) == exp["accepted_rows"]
    assert sorted(rejected, key=lambda r: (r["message"], r["obx"])) == exp["rejections"]
    assert [r["observation_day"] for r in rows] == [54, 54, 85]
    blob = json.dumps(rows)
    assert "LOCAL" not in blob and "TESTSURNAME" not in blob and "2025" not in blob
    assert all(r["patient_id"].startswith("VIGIL-") for r in rows)


def test_hl7_pseudonym_is_deterministic_secret_dependent_and_needs_a_real_secret():
    a, b = hmac_pseudonym(b"0123456789abcdef-A"), hmac_pseudonym(b"0123456789abcdef-B")
    assert a("LOCAL-1") == a("LOCAL-1") != b("LOCAL-1")
    for bad in (b"", b"short"):
        with pytest.raises(ValueError):
            hmac_pseudonym(bad)


def test_hl7_rows_pass_the_same_ingestion_as_any_other_file(tmp_path):
    pseud = hmac_pseudonym(b"0123456789abcdef-test-only")
    rows, _ = convert_messages(_messages(), _index(), pseud)
    pid = pseud("LOCAL-1")
    d = tmp_path / "h"
    d.mkdir()
    for name in ("patients", "baseline_profile", "clinical_outcomes"):
        src = (FIX / f"{name}.csv").read_text(encoding="utf-8").splitlines()
        row = next(ln for ln in src[1:] if ln.startswith("PTR-040,"))
        (d / f"{name}.csv").write_text(
            src[0] + "\n" + row.replace("PTR-040", pid, 1) + "\n", encoding="utf-8"
        )
    with (d / "lab_observations.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    res = ingest_path(d)
    assert res.report["rejections"] == []
    assert len(res.tables["lab_observations"]) == 3


# --------------------------------------------------------------------------------------------- documentation agrees with code
DICTIONARY = (ROOT / "docs" / "clinical" / "PARTNERSHIP" / "DATA-DICTIONARY.md").read_text(
    encoding="utf-8"
)
SOURCE = (ROOT / "packages" / "contracts" / "contracts" / "partner" / "ingest.py").read_text(
    encoding="utf-8"
)


@pytest.mark.parametrize("code", FATAL_CODES + ROW_CODES)
def test_every_rejection_code_is_documented_and_used(code):
    assert f"`{code}`" in DICTIONARY
    assert SOURCE.count(f'"{code}"') >= 2  # the constant and at least one use


def test_every_code_the_code_emits_is_declared(csv_result):
    """The messy dataset exercises every per-row code, and no code outside the declared set is emitted."""
    assert set(csv_result.report["rejection_codes"]) == set(ROW_CODES)


def test_every_column_of_every_table_is_in_the_data_dictionary():
    for table in TABLES:
        for col in table.column_names:
            assert f"`{col}`" in DICTIONARY, f"{table.name}.{col} missing from DATA-DICTIONARY.md"


# --------------------------------------------------------------------------------------------- command line
def _cli():
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "run_partner_ingest", ROOT / "scripts" / "run_partner_ingest.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_cli_exit_codes_and_outputs(tmp_path, capsys):
    cli = _cli()
    assert (
        cli.main([str(FIX), "--out", str(tmp_path / "a")]) == 2
    )  # messy set exceeds the proposed stop rate
    assert (tmp_path / "a" / "clean_patients.csv").exists() and (
        tmp_path / "a" / "ingest_report.md"
    ).exists()
    assert cli.main([str(FIX), "--out", str(tmp_path / "b"), "--stop-rate", "0.5"]) == 0
    bad = _copy(tmp_path)
    _edit(bad / "patients.csv", _add_column("patient_name", "Jane Doe"))
    assert cli.main([str(bad), "--out", str(tmp_path / "c")]) == 1
    assert not list((tmp_path / "c").glob("clean_*.csv"))  # a refusal writes no data
    assert "Jane Doe" not in (tmp_path / "c" / "ingest_report.md").read_text(encoding="utf-8")
