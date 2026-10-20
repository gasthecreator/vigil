"""Partner dataset ingestion with per-row rejection reports.

Policy (fixed, tested):
* A file that contains a direct identifier (identifier-like column, unknown column, absolute date, e-mail address,
  phone number) is refused *as a whole*; nothing is parsed further and the report names columns and row numbers only.
* Otherwise every row is validated on its own; a failing row is rejected, never repaired. Conventional spellings
  (case, whitespace, ``yes``/``Y``/``1``) are normalised and counted, never guessed.
* Conflicting duplicates are all rejected (we do not pick a winner); exact duplicates are dropped and counted.
* Implausible laboratory values are *flagged and withheld* from the clean table pending chart audit.
* The report never contains a raw cell value except a number that passed type validation (and the pseudonymous id).
  A free-text cell that fails validation could hold a name; it is reported as a code, never as text.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from vigil_core.analytes import ANALYTE_REGISTRY, check_plausibility
from vigil_core.units import (
    IncompatibleUnitsError,
    InvalidUnitError,
    convert_units,
    normalize_unit_string,
)

from .tables import (
    FORBIDDEN_COLUMN_FRAGMENTS,
    LAB_OBSERVATIONS,
    TABLE_BY_NAME,
    TABLES,
    Col,
    Table,
)

TOOL_VERSION = "1"
NULL_TOKENS = {"", "na", "n/a", "null", "none", "nan", "-", "--", "."}
_TRUE = {"true", "t", "yes", "y", "1"}
_FALSE = {"false", "f", "no", "n", "0"}

_ISO_DATE = re.compile(r"(?<!\d)\d{4}-\d{2}-\d{2}(?!\d)")
_SLASH_DATE = re.compile(r"(?<!\d)\d{1,2}[/.]\d{1,2}[/.]\d{2,4}(?!\d)")
_EMAIL = re.compile(r"[^\s@]+@[^\s@]+\.[^\s@]+")
_PHONE = re.compile(r"(?<![\w.])\+?\d[\d\s\-()]{8,}\d(?![\w])")
_PSEUDONYM = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_\-]{0,63}$")
_IDENTIFIER_VALUE_PATTERNS = (
    ("ISO_DATE", _ISO_DATE),
    ("SLASH_DATE", _SLASH_DATE),
    ("EMAIL", _EMAIL),
    ("PHONE", _PHONE),
)

LAB_FLAG_REASON = "IMPLAUSIBLE_VALUE"

FATAL_CODES = (
    "MISSING_TABLE",
    "MISSING_COLUMN",
    "IDENTIFIER_COLUMN",
    "UNKNOWN_COLUMN",
    "IDENTIFIER_VALUE",
)
ROW_CODES = (
    "MISSING_REQUIRED",
    "BAD_TYPE",
    "BAD_ENUM",
    "OUT_OF_RANGE",
    "CONFLICTING_DUPLICATE",
    "ORPHAN_FK",
    "PARENT_REJECTED",
    "INCONSISTENT",
    "UNKNOWN_ANALYTE",
    "NAME_CODE_MISMATCH",
    "UNKNOWN_UNIT",
    "INCOMPATIBLE_UNIT",
)


class PartnerIngestAborted(Exception):
    """The input was refused as a whole. ``report`` is safe to hand back to the partner."""

    def __init__(self, report: Dict[str, Any]):
        super().__init__("; ".join(sorted({f["code"] for f in report["fatal"]})))
        self.report = report


@dataclass
class Rejection:
    table: str
    row: int  # 1-based data row (header excluded)
    code: str
    field: str = ""
    patient_id: str = ""
    value: Optional[float] = None  # only ever a validated number

    def as_dict(self) -> Dict[str, Any]:
        d: Dict[str, Any] = {
            "table": self.table,
            "row": self.row,
            "code": self.code,
            "field": self.field,
        }
        if self.patient_id:
            d["patient_id"] = self.patient_id
        if self.value is not None:
            d["value"] = self.value
        return d


@dataclass
class IngestResult:
    tables: Dict[str, List[Dict[str, Any]]]
    flagged_labs: List[Dict[str, Any]]
    rejections: List[Rejection]
    report: Dict[str, Any]

    @property
    def cohort_ids(self) -> List[str]:
        return list(self.report["cohort"]["complete_case_ids"])


# --------------------------------------------------------------------------------------------------- readers
RawRows = List[Tuple[int, Dict[str, Any]]]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv_dir(directory: Path) -> Tuple[Dict[str, RawRows], List[Dict[str, Any]]]:
    directory = Path(directory)
    raw: Dict[str, RawRows] = {}
    inputs: List[Dict[str, Any]] = []
    for table in TABLES:
        path = directory / f"{table.name}.csv"
        if not path.exists():
            continue
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            header = [h.strip() for h in (reader.fieldnames or [])]
            rows: RawRows = []
            for i, rec in enumerate(reader, start=1):
                rows.append((i, {k.strip(): v for k, v in rec.items() if k is not None}))
        raw[table.name] = rows
        raw[table.name + "#header"] = [(0, {"header": header})]  # type: ignore[list-item]
        inputs.append(
            {"table": table.name, "source": path.name, "rows": len(rows), "sha256": _sha256(path)}
        )
    return raw, inputs


def read_xlsx(path: Path) -> Tuple[Dict[str, RawRows], List[Dict[str, Any]]]:
    import datetime as _dt

    from openpyxl import load_workbook

    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    raw: Dict[str, RawRows] = {}
    inputs: List[Dict[str, Any]] = []
    digest = _sha256(path)
    for ws in wb.worksheets:
        name = ws.title.strip().lower()
        if name not in TABLE_BY_NAME:
            continue
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, ())]
        rows: RawRows = []
        for i, values in enumerate(it, start=1):
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            rec: Dict[str, Any] = {}
            for h, v in zip(header, values):
                if isinstance(v, (_dt.datetime, _dt.date)):
                    v = (
                        v.isoformat()
                    )  # surfaces as an absolute date and is refused by the identifier scan
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                rec[h] = "" if v is None else str(v)
            rows.append((i, rec))
        raw[name] = rows
        raw[name + "#header"] = [(0, {"header": header})]  # type: ignore[list-item]
        inputs.append(
            {
                "table": name,
                "source": f"{path.name}#{ws.title}",
                "rows": len(rows),
                "sha256": digest,
            }
        )
    wb.close()
    return raw, inputs


# --------------------------------------------------------------------------------------------------- parsing
def _is_null(s: str) -> bool:
    return s.strip().lower() in NULL_TOKENS


def _parse_cell(col: Col, text: str) -> Tuple[Any, Optional[str], bool]:
    """Return ``(value, error_code, normalised)``. ``value`` is None for null cells."""
    raw = text
    s = text.strip()
    if _is_null(s):
        return None, None, False
    if col.kind == "id":
        if not _PSEUDONYM.match(s):
            return None, "BAD_TYPE", False
        return s, None, False
    if col.kind == "str":
        return s, None, s != raw
    if col.kind == "enum":
        canon = re.sub(r"[\s\-]+", "_", s.upper())
        if canon not in col.enum:
            return None, "BAD_ENUM", False
        return canon, None, canon != raw
    if col.kind == "bool":
        low = s.lower()
        if low in _TRUE:
            return True, None, s not in ("true", "false")
        if low in _FALSE:
            return False, None, s not in ("true", "false")
        return None, "BAD_TYPE", False
    if col.kind in ("int", "float"):
        try:
            num = float(s)
        except ValueError:
            return None, "BAD_TYPE", False
        if not math.isfinite(num):
            return None, "BAD_TYPE", False
        if col.kind == "int":
            if num != int(num):
                return None, "BAD_TYPE", False
            return int(num), None, s != str(int(num))
        return num, None, s != raw
    return None, "BAD_TYPE", False  # pragma: no cover


def _range_error(col: Col, value: Any) -> bool:
    if col.lo is not None and value < col.lo:
        return True
    return col.hi is not None and value > col.hi


# --------------------------------------------------------------------------------------------------- scans
def _identifier_tokens(column: str) -> List[str]:
    return [t for t in re.split(r"[^a-z0-9]+", column.lower()) if t]


def _fatal_scan(raw: Dict[str, RawRows]) -> List[Dict[str, Any]]:
    fatal: List[Dict[str, Any]] = []
    for table in TABLES:
        if table.name not in raw:
            fatal.append({"code": "MISSING_TABLE", "table": table.name})
            continue
        header = raw[table.name + "#header"][0][1]["header"]
        known = set(table.column_names)
        for col in table.columns:
            if col.required and col.name not in header:
                fatal.append({"code": "MISSING_COLUMN", "table": table.name, "column": col.name})
        for h in header:
            if h in known:
                continue
            ident = any(
                tok in FORBIDDEN_COLUMN_FRAGMENTS
                or any(tok.startswith(f) for f in ("phone", "email", "addr"))
                for tok in _identifier_tokens(h)
            )
            fatal.append(
                {
                    "code": "IDENTIFIER_COLUMN" if ident else "UNKNOWN_COLUMN",
                    "table": table.name,
                    "column_position": header.index(h) + 1,
                }
            )
        for row_no, rec in raw[table.name]:
            for colname, cell in rec.items():
                for kind, pat in _IDENTIFIER_VALUE_PATTERNS:
                    if isinstance(cell, str) and pat.search(cell):
                        fatal.append(
                            {
                                "code": "IDENTIFIER_VALUE",
                                "kind": kind,
                                "table": table.name,
                                "row": row_no,
                                "field": colname,
                            }
                        )
                        break
    return fatal


# --------------------------------------------------------------------------------------------------- table pass
def _validate_table(
    table: Table, rows: RawRows, rejections: List[Rejection], counters: Dict[str, Any]
) -> List[Tuple[int, Dict[str, Any]]]:
    typed: List[Tuple[int, Dict[str, Any]]] = []
    for row_no, rec in rows:
        out: Dict[str, Any] = {}
        pid = (rec.get("patient_id") or "").strip()
        pid_ok = bool(_PSEUDONYM.match(pid))
        bad = False
        for col in table.columns:
            text = rec.get(col.name, "")
            value, err, normalised = _parse_cell(col, text if isinstance(text, str) else str(text))
            if normalised:
                counters["normalised_cells"] += 1
            if err:
                rejections.append(
                    Rejection(table.name, row_no, err, col.name, pid if pid_ok else "")
                )
                bad = True
                continue
            if value is None:
                if col.required:
                    rejections.append(
                        Rejection(
                            table.name, row_no, "MISSING_REQUIRED", col.name, pid if pid_ok else ""
                        )
                    )
                    bad = True
                out[col.name] = None
                continue
            if col.kind in ("int", "float") and _range_error(col, value):
                rejections.append(
                    Rejection(
                        table.name,
                        row_no,
                        "OUT_OF_RANGE",
                        col.name,
                        pid if pid_ok else "",
                        float(value),
                    )
                )
                bad = True
                continue
            out[col.name] = value
        if not bad:
            typed.append((row_no, out))
    return typed


def _dedupe(
    table: Table,
    rows: List[Tuple[int, Dict[str, Any]]],
    rejections: List[Rejection],
    counters: Dict[str, Any],
) -> List[Tuple[int, Dict[str, Any]]]:
    """Exact duplicates dropped (counted); conflicting duplicates all rejected."""
    groups: Dict[Tuple[Any, ...], List[Tuple[int, Dict[str, Any]]]] = {}
    for item in rows:
        groups.setdefault(tuple(item[1][k] for k in table.key), []).append(item)
    keep: List[Tuple[int, Dict[str, Any]]] = []
    for items in groups.values():
        if len(items) == 1:
            keep.append(items[0])
            continue
        payloads = {json.dumps(i[1], sort_keys=True, default=str) for i in items}
        if len(payloads) == 1:
            keep.append(items[0])
            counters["duplicates_dropped"] += len(items) - 1
            counters["dups"][table.name] = counters["dups"].get(table.name, 0) + len(items) - 1
        else:
            for row_no, rec in items:
                rejections.append(
                    Rejection(
                        table.name,
                        row_no,
                        "CONFLICTING_DUPLICATE",
                        ",".join(table.key),
                        rec["patient_id"],
                    )
                )
    return sorted(keep, key=lambda x: x[0])


def ingest(
    raw: Dict[str, RawRows], inputs: List[Dict[str, Any]], stop_rate: float = 0.10
) -> IngestResult:
    """Validate a raw extraction. ``stop_rate`` is the rejection-rate stop threshold (PREREG PR-07, once ratified)."""
    fatal = _fatal_scan(raw)
    base_report: Dict[str, Any] = {
        "tool": "contracts.partner.ingest",
        "tool_version": TOOL_VERSION,
        "inputs": inputs,
    }
    if fatal:
        raise PartnerIngestAborted({**base_report, "aborted": True, "fatal": fatal})

    rejections: List[Rejection] = []
    counters: Dict[str, Any] = {"normalised_cells": 0, "duplicates_dropped": 0, "dups": {}}
    per_table: Dict[str, Dict[str, int]] = {}
    typed: Dict[str, List[Tuple[int, Dict[str, Any]]]] = {}
    for table in TABLES:
        t = _validate_table(table, raw[table.name], rejections, counters)
        typed[table.name] = (
            _dedupe(table, t, rejections, counters) if table is not LAB_OBSERVATIONS else t
        )
        per_table[table.name] = {"rows_in": len(raw[table.name])}

    # Patient timing must be explicit; total follow-up is not a substitute for a death date.
    consistent: List[Tuple[int, Dict[str, Any]]] = []
    for row_no, rec in typed["patients"]:
        deceased = rec["vital_status"] in {"DEAD_DISEASE", "DEAD_OTHER"}
        death_day = rec["days_to_death"]
        if rec["index_day"] > rec["days_to_last_followup"]:
            rejections.append(
                Rejection("patients", row_no, "INCONSISTENT", "index_day", rec["patient_id"])
            )
        elif deceased and death_day is None:
            rejections.append(
                Rejection("patients", row_no, "MISSING_REQUIRED", "days_to_death", rec["patient_id"])
            )
        elif not deceased and death_day is not None:
            rejections.append(
                Rejection("patients", row_no, "INCONSISTENT", "days_to_death", rec["patient_id"])
            )
        elif death_day is not None and not rec["index_day"] <= death_day <= rec["days_to_last_followup"]:
            rejections.append(
                Rejection("patients", row_no, "INCONSISTENT", "days_to_death", rec["patient_id"])
            )
        else:
            consistent.append((row_no, rec))
    typed["patients"] = consistent
    # parent rows: the patients table decides who exists
    patient_ids = {r["patient_id"] for _, r in typed["patients"]}
    all_listed_ids = {(rec.get("patient_id") or "").strip() for _, rec in raw["patients"]}

    def _fk(table_name: str) -> List[Tuple[int, Dict[str, Any]]]:
        kept = []
        for row_no, rec in typed[table_name]:
            pid = rec["patient_id"]
            if pid in patient_ids:
                kept.append((row_no, rec))
            elif pid in all_listed_ids:
                rejections.append(
                    Rejection(table_name, row_no, "PARENT_REJECTED", "patient_id", pid)
                )
            else:
                rejections.append(Rejection(table_name, row_no, "ORPHAN_FK", "patient_id", pid))
        return kept

    typed["baseline_profile"] = _fk("baseline_profile")
    typed["clinical_outcomes"] = _fk("clinical_outcomes")
    labs_in = _fk("lab_observations")

    # outcomes: conditional requirements and consistency with follow-up
    followup = {r["patient_id"]: r["days_to_last_followup"] for _, r in typed["patients"]}
    outcomes: List[Tuple[int, Dict[str, Any]]] = []
    for row_no, rec in typed["clinical_outcomes"]:
        pid = rec["patient_id"]
        sites = [
            rec[k]
            for k in ("site_bone", "site_liver", "site_lung", "site_brain", "site_locoregional")
        ]
        if rec["recurrence_confirmed"]:
            if rec["recurrence_day"] is None:
                rejections.append(
                    Rejection(
                        "clinical_outcomes", row_no, "MISSING_REQUIRED", "recurrence_day", pid
                    )
                )
                continue
            if rec["recurrence_day"] > followup[pid]:
                rejections.append(
                    Rejection("clinical_outcomes", row_no, "INCONSISTENT", "recurrence_day", pid)
                )
                continue
        elif rec["recurrence_day"] is not None or any(sites):
            rejections.append(
                Rejection("clinical_outcomes", row_no, "INCONSISTENT", "recurrence_confirmed", pid)
            )
            continue
        outcomes.append((row_no, rec))
    typed["clinical_outcomes"] = outcomes

    # laboratory rows: analyte, unit, canonicalisation, plausibility, conflicting duplicates
    loinc_to_code = {a.loinc_code: a.code for a in ANALYTE_REGISTRY.values()}
    labs: List[Tuple[int, Dict[str, Any]]] = []
    flagged: List[Dict[str, Any]] = []
    for row_no, rec in labs_in:
        pid = rec["patient_id"]
        code = loinc_to_code.get(rec["analyte_code"])
        if code is None:
            rejections.append(
                Rejection("lab_observations", row_no, "UNKNOWN_ANALYTE", "analyte_code", pid)
            )
            continue
        if rec["analyte_name"].strip().lower() != code:
            rejections.append(
                Rejection("lab_observations", row_no, "NAME_CODE_MISMATCH", "analyte_name", pid)
            )
            continue
        analyte = ANALYTE_REGISTRY[code]
        try:
            normalize_unit_string(rec["unit_ucum"])
        except InvalidUnitError:
            rejections.append(
                Rejection("lab_observations", row_no, "UNKNOWN_UNIT", "unit_ucum", pid)
            )
            continue
        try:
            canonical = convert_units(
                rec["value_numeric"], rec["unit_ucum"], analyte.canonical_unit, analyte_code=code
            )
        except IncompatibleUnitsError:
            rejections.append(
                Rejection("lab_observations", row_no, "INCOMPATIBLE_UNIT", "unit_ucum", pid)
            )
            continue
        row = {
            **rec,
            "analyte": code,
            "value_canonical": canonical,
            "unit_canonical": analyte.canonical_unit,
        }
        labs.append((row_no, row))
    groups: Dict[Tuple[Any, ...], List[Tuple[int, Dict[str, Any]]]] = {}
    for item in labs:
        groups.setdefault(
            (item[1]["patient_id"], item[1]["observation_day"], item[1]["analyte"]), []
        ).append(item)
    labs_clean: List[Tuple[int, Dict[str, Any]]] = []
    for items in groups.values():
        if len(items) > 1:
            vals = {round(i[1]["value_canonical"], 9) for i in items}
            if len(vals) == 1:
                counters["duplicates_dropped"] += len(items) - 1
                counters["dups"]["lab_observations"] = (
                    counters["dups"].get("lab_observations", 0) + len(items) - 1
                )
                items = items[:1]
            else:
                for row_no, rec in items:
                    rejections.append(
                        Rejection(
                            "lab_observations",
                            row_no,
                            "CONFLICTING_DUPLICATE",
                            "patient_id,observation_day,analyte_code",
                            rec["patient_id"],
                        )
                    )
                continue
        labs_clean.extend(items)
    final_labs: List[Dict[str, Any]] = []
    for row_no, rec in sorted(labs_clean, key=lambda x: x[0]):
        res = check_plausibility(rec["analyte"], rec["value_numeric"], rec["unit_ucum"])
        if not res.is_valid:
            flagged.append(
                {**rec, "source_row": row_no, "flag": LAB_FLAG_REASON, "status": res.status.value}
            )
        else:
            final_labs.append({**rec, "source_row": row_no})

    clean = {
        "patients": [{**r, "source_row": n} for n, r in typed["patients"]],
        "baseline_profile": [{**r, "source_row": n} for n, r in typed["baseline_profile"]],
        "lab_observations": final_labs,
        "clinical_outcomes": [{**r, "source_row": n} for n, r in typed["clinical_outcomes"]],
    }

    codes: Dict[str, int] = {}
    for r in rejections:
        codes[r.code] = codes.get(r.code, 0) + 1
    rows_in_total = sum(per_table[t.name]["rows_in"] for t in TABLES)
    for t in TABLES:
        rej = len({r.row for r in rejections if r.table == t.name})
        per_table[t.name] = {
            "rows_in": per_table[t.name]["rows_in"],
            "rows_clean": len(clean[t.name]),
            "rows_rejected": rej,
            "duplicates_dropped": counters["dups"].get(t.name, 0),
            "rows_flagged": len(flagged) if t is LAB_OBSERVATIONS else 0,
        }
    have_base = {r["patient_id"] for r in clean["baseline_profile"]}
    have_out = {r["patient_id"] for r in clean["clinical_outcomes"]}
    complete = sorted({r["patient_id"] for r in clean["patients"]} & have_base & have_out)
    rate = (
        sum(per_table[t.name]["rows_rejected"] for t in TABLES) / rows_in_total
        if rows_in_total
        else 0.0
    )
    report = {
        **base_report,
        "aborted": False,
        "fatal": [],
        "per_table": per_table,
        "rejection_codes": dict(sorted(codes.items())),
        "rejection_rate": rate,
        "stop_rate": stop_rate,
        "exceeds_stop_rate": rate > stop_rate,
        "duplicates_dropped": counters["duplicates_dropped"],
        "normalised_cells": counters["normalised_cells"],
        "cohort": {
            "patients_listed": len(all_listed_ids),
            "complete_case": len(complete),
            "complete_case_ids": complete,
        },
        "rejections": [
            r.as_dict() for r in sorted(rejections, key=lambda r: (r.table, r.row, r.code, r.field))
        ],
        "flagged_labs": [
            {
                "row": f["source_row"],
                "patient_id": f["patient_id"],
                "analyte": f["analyte"],
                "value": f["value_numeric"],
                "status": f["status"],
            }
            for f in flagged
        ],
    }
    return IngestResult(clean, flagged, rejections, report)


def ingest_path(path: Path, stop_rate: float = 0.10) -> IngestResult:
    path = Path(path)
    raw, inputs = (
        read_xlsx(path)
        if path.is_file() and path.suffix.lower() in (".xlsx", ".xlsm")
        else read_csv_dir(path)
    )
    return ingest(raw, inputs, stop_rate)


def render_markdown(report: Dict[str, Any], title: str = "Partner ingestion report") -> str:
    L = [f"# {title}", ""]
    if report.get("aborted"):
        L += [
            "**The input was refused as a whole.** No row was parsed.",
            "",
            "| code | table | where |",
            "|---|---|---|",
        ]
        for f in report["fatal"]:
            where = f.get("column") or (
                f"column #{f['column_position']}" if "column_position" in f else ""
            )
            if "row" in f:
                where = f"row {f['row']}, column `{f['field']}` ({f['kind']})"
            L.append(f"| {f['code']} | {f.get('table', '')} | {where} |")
        return "\n".join(L) + "\n"
    L += ["| table | rows in | clean | rejected | flagged |", "|---|---:|---:|---:|---:|"]
    for t, v in report["per_table"].items():
        L.append(
            f"| {t} | {v['rows_in']} | {v['rows_clean']} | {v['rows_rejected']} | {v['rows_flagged']} |"
        )
    L += [
        "",
        f"Rejection rate {report['rejection_rate']:.3f} against stop rate {report['stop_rate']:.3f}: "
        + (
            "**EXCEEDS the stop rate. Stop and review the extraction with the partner.**"
            if report["exceeds_stop_rate"]
            else "within the stop rate."
        ),
        "",
    ]
    L += ["| rejection code | rows |", "|---|---:|"] + [
        f"| {c} | {n} |" for c, n in report["rejection_codes"].items()
    ]
    L += [
        "",
        f"Exact duplicates dropped: {report['duplicates_dropped']}. Cells normalised (case, spelling): {report['normalised_cells']}. Complete-case patients (patient + baseline + outcome): {report['cohort']['complete_case']} of {report['cohort']['patients_listed']}.",
        "",
    ]
    return "\n".join(L)
